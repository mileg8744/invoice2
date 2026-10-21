"""Invoice PPTX built from the official A4 PDF so layout stays identical."""

from io import BytesIO
from pathlib import Path
import tempfile

from pptx import Presentation
from pptx.util import Emu, Mm

from pdf_invoice import generate_invoice_pdf


def generate_invoice_ppt(invoice, output_path: Path, show_bank: bool = True) -> Path:
    output_path = Path(output_path)
    output_path.parent.mkdir(parents=True, exist_ok=True)
    with tempfile.TemporaryDirectory() as tmp:
        pdf_path = Path(tmp) / "invoice.pdf"
        generate_invoice_pdf(invoice, pdf_path, show_bank=show_bank)
        _pdf_pages_to_pptx(pdf_path, output_path)
    return output_path


def pdf_bytes_to_pptx(pdf_bytes: bytes, output_path: Path) -> Path:
    output_path = Path(output_path)
    output_path.parent.mkdir(parents=True, exist_ok=True)
    if not pdf_bytes:
        raise ValueError("empty pdf")
    with tempfile.TemporaryDirectory() as tmp:
        pdf_path = Path(tmp) / "source.pdf"
        pdf_path.write_bytes(pdf_bytes)
        _pdf_pages_to_pptx(pdf_path, output_path)
    return output_path


def _pdf_pages_to_pptx(pdf_path: Path, output_path: Path) -> None:
    import pypdfium2 as pdfium

    prs = Presentation()
    prs.slide_width = Mm(210)
    prs.slide_height = Mm(297)
    blank = prs.slide_layouts[6]
    pdf = pdfium.PdfDocument(str(pdf_path))
    try:
        for index in range(len(pdf)):
            page = pdf[index]
            bitmap = page.render(scale=1.6)
            image = bitmap.to_pil().convert("RGB")
            buf = BytesIO()
            image.save(buf, format="JPEG", quality=72, optimize=True)
            buf.seek(0)
            slide = prs.slides.add_slide(blank)
            slide.shapes.add_picture(buf, Emu(0), Emu(0), width=prs.slide_width, height=prs.slide_height)
    finally:
        pdf.close()
    prs.save(str(output_path))
