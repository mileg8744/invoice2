from datetime import date, datetime
from decimal import Decimal, ROUND_HALF_UP
from functools import wraps
from io import BytesIO, StringIO
from pathlib import Path
import csv
import json
import logging
import re

from flask import (
    Flask,
    Response,
    flash,
    g,
    redirect,
    render_template,
    request,
    send_file,
    session,
    url_for,
)
from openpyxl import Workbook, load_workbook
from sqlalchemy import create_engine, func, inspect, or_
from sqlalchemy.orm import object_session, scoped_session, sessionmaker
from werkzeug.security import check_password_hash, generate_password_hash
from PIL import Image
from werkzeug.utils import secure_filename

from config import Config, INSTANCE_DIR, PDF_DIR, STAMP_PATH, STATIC_DIR, UPLOAD_DIR
from i18n import translate
from mailer import (
    draft_mail_content,
    hq_copy_emails,
    mail_placeholder_help,
    parse_emails,
    send_invoice_email,
    send_test_email,
)
from smtp_mail import smtp_configured
from models import AuditLog, Base, Invoice, InvoiceItem, Subsidiary, Training, User, utcnow
from pdf_invoice import generate_invoice_pdf
from regions import REGIONS, apply_region, resolve_region
from entity_types import ENTITY_TYPES, entity_type_label, resolve_entity_type
from seed import seed_if_empty
from hq_settings import load_hq_settings, save_hq_settings
from corporation_master import (
    MASTER_JSON,
    contact_emails,
    emails_from_text,
    emails_to_text,
    ensure_login_user,
    freeze_issued_parties,
    hr_emails,
    import_master_json,
    invoice_party,
    mark_duplicate_codes,
    snapshot_invoice_party,
    split_email_roles,
)

STATUS_ORDER = ["draft", "issued", "acknowledged", "paid"]


def create_app():
    app = Flask(__name__)
    app.config.from_object(Config)
    logging.getLogger("pgu.mail").setLevel(logging.INFO)
    INSTANCE_DIR.mkdir(parents=True, exist_ok=True)
    UPLOAD_DIR.mkdir(parents=True, exist_ok=True)
    PDF_DIR.mkdir(parents=True, exist_ok=True)
    (STATIC_DIR / "img").mkdir(parents=True, exist_ok=True)
    load_hq_settings()

    db_uri = Config.SQLALCHEMY_DATABASE_URI
    engine_kwargs = {"echo": False, "future": True}
    if db_uri.startswith("postgresql"):
        engine_kwargs.update(pool_pre_ping=True)
    elif db_uri.startswith("sqlite"):
        engine_kwargs["connect_args"] = {"check_same_thread": False}
    engine = create_engine(db_uri, **engine_kwargs)
    Session = scoped_session(sessionmaker(bind=engine, autoflush=False, autocommit=False, expire_on_commit=False))
    Base.metadata.create_all(engine)
    with engine.begin() as conn:
        inspector = inspect(conn)
        table_names = set(inspector.get_table_names())
        if "invoice_items" in table_names:
            item_cols = {c["name"] for c in inspector.get_columns("invoice_items")}
            for name, ddl in [
                ("training_id", "ALTER TABLE invoice_items ADD COLUMN training_id INTEGER"),
                ("learners", "ALTER TABLE invoice_items ADD COLUMN learners TEXT DEFAULT ''"),
                ("instructor", "ALTER TABLE invoice_items ADD COLUMN instructor VARCHAR(100) DEFAULT ''"),
                ("period_start", "ALTER TABLE invoice_items ADD COLUMN period_start DATE"),
                ("period_end", "ALTER TABLE invoice_items ADD COLUMN period_end DATE"),
                ("learner_roster", "ALTER TABLE invoice_items ADD COLUMN learner_roster TEXT DEFAULT ''"),
            ]:
                if name not in item_cols:
                    conn.exec_driver_sql(ddl)
        if "trainings" in table_names:
            train_cols = {c["name"] for c in inspector.get_columns("trainings")}
            if "learners" not in train_cols:
                conn.exec_driver_sql("ALTER TABLE trainings ADD COLUMN learners TEXT DEFAULT ''")
            if "learner_count" not in train_cols:
                conn.exec_driver_sql("ALTER TABLE trainings ADD COLUMN learner_count INTEGER DEFAULT 0")
            if "title_en" not in train_cols:
                conn.exec_driver_sql("ALTER TABLE trainings ADD COLUMN title_en VARCHAR(300) DEFAULT ''")
                conn.exec_driver_sql("UPDATE trainings SET title_en = title WHERE title_en IS NULL OR title_en = ''")
        if "subsidiaries" in table_names:
            sub_cols = {c["name"] for c in inspector.get_columns("subsidiaries")}
            if "notes" not in sub_cols:
                conn.exec_driver_sql("ALTER TABLE subsidiaries ADD COLUMN notes TEXT DEFAULT ''")
            if "entity_type" not in sub_cols:
                conn.exec_driver_sql("ALTER TABLE subsidiaries ADD COLUMN entity_type VARCHAR(80) DEFAULT ''")
            if "list_no" not in sub_cols:
                conn.exec_driver_sql("ALTER TABLE subsidiaries ADD COLUMN list_no INTEGER")
            if "company" not in sub_cols:
                conn.exec_driver_sql("ALTER TABLE subsidiaries ADD COLUMN company VARCHAR(80) DEFAULT ''")
            if "hr_emails" not in sub_cols:
                conn.exec_driver_sql("ALTER TABLE subsidiaries ADD COLUMN hr_emails TEXT DEFAULT ''")
            if "needs_review" not in sub_cols:
                conn.exec_driver_sql("ALTER TABLE subsidiaries ADD COLUMN needs_review BOOLEAN DEFAULT FALSE")
            try:
                for uq in inspector.get_unique_constraints("subsidiaries"):
                    cols = list(uq.get("column_names") or [])
                    if cols == ["code"] and uq.get("name"):
                        conn.exec_driver_sql(f'ALTER TABLE subsidiaries DROP CONSTRAINT IF EXISTS "{uq["name"]}"')
                for idx in inspector.get_indexes("subsidiaries"):
                    cols = list(idx.get("column_names") or [])
                    if idx.get("unique") and cols == ["code"] and idx.get("name"):
                        conn.exec_driver_sql(f'DROP INDEX IF EXISTS "{idx["name"]}"')
                        conn.exec_driver_sql(
                            "CREATE INDEX IF NOT EXISTS ix_subsidiaries_code ON subsidiaries (code)"
                        )
            except Exception:
                pass
        if "invoices" in table_names:
            inv_cols = {c["name"] for c in inspector.get_columns("invoices")}
            if "billed_name" not in inv_cols:
                conn.exec_driver_sql("ALTER TABLE invoices ADD COLUMN billed_name VARCHAR(200) DEFAULT ''")
            if "billed_address" not in inv_cols:
                conn.exec_driver_sql("ALTER TABLE invoices ADD COLUMN billed_address TEXT DEFAULT ''")
            if "billed_code" not in inv_cols:
                conn.exec_driver_sql("ALTER TABLE invoices ADD COLUMN billed_code VARCHAR(20) DEFAULT ''")

    with Session() as db:
        seed_if_empty(db)
        for sub in db.query(Subsidiary).all():
            apply_region(sub)
        freeze_issued_parties(db)
        mark_duplicate_codes(db)
        db.commit()

    def db_session():
        return Session()

    @app.teardown_appcontext
    def shutdown_session(exception=None):
        Session.remove()

    def t(key):
        return translate(session.get("lang", "ko"), key)

    def current_user():
        uid = session.get("user_id")
        if not uid:
            return None
        if getattr(g, "_user", None) and g._user.id == uid:
            return g._user
        db = Session()
        user = db.get(User, uid)
        g._user = user
        return user

    def log_action(db, entity_type, entity_id, action, details=""):
        user = current_user()
        db.add(
            AuditLog(
                entity_type=entity_type,
                entity_id=str(entity_id or ""),
                action=action,
                username=user.username if user else "",
                details=details,
            )
        )

    def login_required(fn):
        @wraps(fn)
        def wrapper(*args, **kwargs):
            if not current_user():
                flash(t("need_login"), "warning")
                return redirect(url_for("login"))
            return fn(*args, **kwargs)

        return wrapper

    def admin_required(fn):
        @wraps(fn)
        def wrapper(*args, **kwargs):
            user = current_user()
            if not user:
                return redirect(url_for("login"))
            if user.role != "admin":
                flash(t("no_permission"), "danger")
                return redirect(url_for("dashboard"))
            return fn(*args, **kwargs)

        return wrapper

    @app.context_processor
    def inject_globals():
        user = current_user()
        lang = session.get("lang", "ko")
        return {
            "t": t,
            "lang": lang,
            "current_user": user,
            "today": date.today().strftime("%Y-%m-%d"),
            "status_label": lambda s: t(f"status_{s}") if s else "",
            "entity_types": ENTITY_TYPES,
            "entity_type_label": entity_type_label,
            "hq_copy_email": Config.HQ_COPY_EMAIL,
            "mail_live": smtp_configured(),
            "mail_from": getattr(Config, "MAIL_FROM", "") or "",
        }

    @app.template_filter("krw")
    def krw_filter(v):
        try:
            return f"{int(v):,}"
        except (TypeError, ValueError):
            return "0"

    @app.template_filter("fx")
    def fx_filter(v):
        try:
            return f"{float(v):,.2f}"
        except (TypeError, ValueError):
            return "0.00"

    @app.template_filter("rate")
    def rate_filter(v):
        try:
            return f"{float(v):,.2f}"
        except (TypeError, ValueError):
            return "0.00"

    @app.template_filter("ymd")
    def ymd_filter(v):
        if not v:
            return "-"
        if hasattr(v, "strftime"):
            return v.strftime("%Y-%m-%d")
        return str(v)[:10]

    @app.template_filter("dt")
    def dt_filter(v):
        if not v:
            return "-"
        if hasattr(v, "strftime"):
            return v.strftime("%Y-%m-%d %H:%M")
        return str(v)[:16]

    @app.before_request
    def before():
        session.setdefault("lang", "ko")
        if request.endpoint in (
            "login",
            "logout",
            "lang",
            "static",
            "change_password",
        ) or (request.endpoint or "").startswith("static"):
            return None
        user = current_user()
        if user and user.must_change_password and request.endpoint != "change_password":
            return redirect(url_for("change_password"))
        return None

    # ---------- helpers ----------
    def half_from_date(d: date) -> int:
        return 1 if d.month <= 6 else 2

    def next_invoice_no(db, year: int, half: int):
        last = (
            db.query(func.max(Invoice.seq))
            .filter(Invoice.year == year, Invoice.half == half)
            .scalar()
        )
        seq = int(last or 0) + 1
        return f"PGU-{year}-{half:02d}-{seq:03d}", seq

    def money_fx(amount_krw: int, rate: float) -> Decimal:
        if not rate:
            return Decimal("0.00")
        val = Decimal(str(amount_krw)) / Decimal(str(rate))
        return val.quantize(Decimal("0.01"), rounding=ROUND_HALF_UP)

    @app.context_processor
    def inject_money():
        return {"money_fx": money_fx}

    def allowed_file(filename: str) -> bool:
        return "." in filename and filename.rsplit(".", 1)[1].lower() in {"csv", "xlsx", "xls"}

    def allowed_stamp(filename: str) -> bool:
        return "." in filename and filename.rsplit(".", 1)[1].lower() in {"png", "jpg", "jpeg", "webp", "gif"}

    def load_table(file_storage):
        name = secure_filename(file_storage.filename or "")
        ext = name.rsplit(".", 1)[-1].lower()
        data = file_storage.read()
        if ext == "csv":
            text = data.decode("utf-8-sig")
            reader = csv.DictReader(text.splitlines())
            return [dict(row) for row in reader]
        bio = BytesIO(data)
        wb = load_workbook(bio, data_only=True)
        ws = wb.active
        rows = list(ws.iter_rows(values_only=True))
        if not rows:
            return []
        headers = [str(h).strip() if h is not None else "" for h in rows[0]]
        out = []
        for row in rows[1:]:
            if not row or all(c is None or str(c).strip() == "" for c in row):
                continue
            item = {}
            for i, h in enumerate(headers):
                if not h:
                    continue
                val = row[i] if i < len(row) else None
                item[h] = val
            out.append(item)
        return out

    def norm_key(row: dict, *keys):
        mapping = {str(k).strip().lower(): v for k, v in row.items()}
        for key in keys:
            k = key.lower()
            if k in mapping and mapping[k] not in (None, ""):
                return mapping[k]
        return ""

    def parse_period(raw, start=None, end=None):
        if start and end:
            return _as_date(start), _as_date(end)
        text = str(raw or "").strip()
        parts = re.split(r"\s*(?:~|to|-|–|—)\s*", text)
        if len(parts) >= 2 and parts[0] and parts[1]:
            return _as_date(parts[0]), _as_date(parts[1])
        d = _as_date(text) if text else date.today()
        return d, d

    def _as_date(value):
        if isinstance(value, datetime):
            return value.date()
        if isinstance(value, date):
            return value
        text = str(value).strip()[:10]
        return datetime.strptime(text, "%Y-%m-%d").date()

    def _as_int(value, default=0):
        if value in (None, ""):
            return default
        try:
            return int(float(str(value).replace(",", "").strip()))
        except ValueError:
            return default

    def _as_rate(value, default=0.0):
        if value in (None, ""):
            return default
        try:
            return round(float(str(value).replace(",", "").strip()), 2)
        except (TypeError, ValueError):
            return default

    def flash_mail_result(mail, success_key, commit=False):
        if not mail.get("ok"):
            err = mail.get("error") or ""
            if err == "no_email":
                flash(t("no_email"), "danger")
            elif err in (
                "smtp_not_configured",
                "smtp_disabled",
                "graph_not_configured",
                "gmail_not_configured",
                "gmail_not_connected",
            ):
                flash(t("smtp_not_configured"), "danger")
            elif err == "smtp_auth_failed":
                flash(t("smtp_auth_failed"), "danger")
            elif err == "mail_body_required":
                flash(t("mail_body_required"), "danger")
            else:
                flash(f"{t('mail_send_fail')} {err}".strip(), "danger")
            return False
        flash(t(success_key), "success")
        to_list = mail.get("to") or mail.get("recipients") or []
        cc_list = mail.get("cc") or []
        if to_list:
            flash(f"{t('mail_to')}: {', '.join(to_list)}", "info")
        if cc_list:
            flash(f"{t('mail_cc')}: {', '.join(cc_list)}", "info")
        if mail.get("via") == "outbox":
            flash(t("email_saved_outbox"), "info")
        return True

    def mail_log_detail(mail):
        to_list = ",".join(mail.get("to") or mail.get("recipients") or [])
        cc_list = ",".join(mail.get("cc") or [])
        via = mail.get("via") or ""
        parts = [f"to={to_list}"]
        if cc_list:
            parts.append(f"cc={cc_list}")
        if via:
            parts.append(f"via={via}")
        return "; ".join(parts)

    def invoice_compose_recipients(row, kind):
        if kind == "hq":
            return hq_copy_emails(), []
        return split_email_roles(contact_emails(row.subsidiary), hr_emails(row.subsidiary))

    def parse_compose_recipients(form, row, kind):
        extra_to = emails_from_text(form.get("mail_to_extra") or "")
        extra_cc = emails_from_text(form.get("mail_cc_extra") or "")
        if form.get("mail_recipients_ready") == "1":
            return split_email_roles(form.getlist("mail_to") + extra_to, form.getlist("mail_cc") + extra_cc)
        to_list, cc_list = invoice_compose_recipients(row, kind)
        return split_email_roles(to_list + extra_to, cc_list + extra_cc)

    def render_mail_compose(kind, rows, action, cancel, hidden=None, fill=True):
        rows = list(rows or [])
        first = rows[0] if rows else None
        fill_one = bool(fill and first and len(rows) == 1)
        draft = draft_mail_content(kind, first if fill_one else None, fill=fill_one)
        to_parts = []
        cc_parts = []
        for row in rows:
            to_list, cc_list = invoice_compose_recipients(row, kind)
            to_parts.extend(to_list)
            cc_parts.extend(cc_list)
        to_choices = list(dict.fromkeys(to_parts))
        cc_choices = [e for e in dict.fromkeys(cc_parts) if e.lower() not in {x.lower() for x in to_choices}]
        return render_template(
            "invoices/compose.html",
            kind=kind,
            rows=rows,
            draft=draft,
            action=action,
            cancel=cancel,
            hidden=hidden or {},
            selected_ids=[r.id for r in rows],
            to_choices=to_choices,
            cc_choices=cc_choices,
            to_text=", ".join(to_choices),
            cc_text=", ".join(cc_choices),
            placeholders=mail_placeholder_help(),
            bulk=len(rows) > 1,
        )

    EMAIL_RE = re.compile(r"[A-Za-z0-9._%+\-]+@[A-Za-z0-9.\-]+\.[A-Za-z]{2,}")
    CODE_RE = re.compile(r"([0-9]{2}[A-Za-z]{2}[0-9]{2})")
    HEADCOUNT_RE = re.compile(r"(\d+)\s*(?:명|people|persons|pax)", re.I)

    def parse_person(text):
        raw = str(text or "").strip()
        email = ""
        found = EMAIL_RE.search(raw)
        if found:
            email = found.group(0)
            raw = (raw[: found.start()] + raw[found.end() :]).strip(" \t,;|")
        code = ""
        found = CODE_RE.search(raw)
        if found:
            code = found.group(1).upper()
            raw = (raw[: found.start()] + raw[found.end() :]).strip(" \t,;|")
        name = raw.strip(" \t,;|:-")
        return code, name, email

    def parse_people_lines(raw):
        people = []
        for line in str(raw or "").replace("\r", "").splitlines():
            text = line.strip()
            if not text:
                continue
            _code, name, email = parse_person(text)
            if name or email:
                people.append({"name": name, "email": email})
        return people

    def dump_roster(people):
        return json.dumps(people or [], ensure_ascii=False)

    def load_roster(item):
        raw = (getattr(item, "learner_roster", None) or "").strip()
        if raw.startswith("["):
            try:
                data = json.loads(raw)
                if isinstance(data, list):
                    out = []
                    for person in data:
                        if not isinstance(person, dict):
                            continue
                        name = str(person.get("name") or "").strip()
                        email = str(person.get("email") or "").strip()
                        if name or email:
                            out.append({"name": name, "email": email})
                    if out:
                        return out
            except json.JSONDecodeError:
                pass
        people = []
        blob = (getattr(item, "learners", None) or getattr(item, "learner_name", None) or "").strip()
        for line in blob.splitlines():
            _code, name, email = parse_person(line)
            if name or email:
                people.append({"name": name, "email": email})
        return people

    def roster_as_text(people):
        lines = []
        for person in people or []:
            name = (person.get("name") or "").strip()
            email = (person.get("email") or "").strip()
            if name and email:
                lines.append(f"{name}, {email}")
            elif name or email:
                lines.append(name or email)
        return "\n".join(lines)

    def invoice_pdf_path(invoice):
        return PDF_DIR / f"{invoice.invoice_no}.pdf"

    def rebuild_pdf(invoice):
        path = invoice_pdf_path(invoice)
        generate_invoice_pdf(invoice, path)
        invoice.pdf_path = str(path)
        return path

    def apply_invoice_totals(invoice):
        if invoice.items:
            invoice.qty = sum(it.qty for it in invoice.items) or invoice.qty
            invoice.amount_krw = sum(it.amount_krw for it in invoice.items)
            if invoice.qty:
                invoice.unit_price_krw = int(round(invoice.amount_krw / invoice.qty))
        else:
            invoice.amount_krw = int(invoice.qty) * int(invoice.unit_price_krw)
        invoice.amount_fx = money_fx(invoice.amount_krw, invoice.exchange_rate)

    def apply_invoice_header_from_items(invoice):
        items = list(invoice.items or [])
        titles = [(it.description or "").strip() for it in items if (it.description or "").strip()]
        if len(titles) == 1:
            invoice.education_name = titles[0]
        elif titles:
            joined = " / ".join(titles)
            invoice.education_name = joined if len(joined) <= 300 else f"{titles[0]} (+{len(titles) - 1})"
        starts = [it.period_start for it in items if it.period_start]
        ends = [it.period_end for it in items if it.period_end]
        if starts:
            invoice.period_start = min(starts)
        if ends:
            invoice.period_end = max(ends)
        invoice.instructor = ", ".join(
            dict.fromkeys(it.instructor for it in items if it.instructor)
        )
        unique_tids = list(dict.fromkeys(it.training_id for it in items if it.training_id))
        invoice.training_id = unique_tids[0] if len(unique_tids) == 1 else None
        apply_invoice_totals(invoice)

    def invoice_item_from_training(training, people):
        qty = billed_qty_for_entity(training, people=people, fallback=max(len(people or []), 1))
        named = [p["name"] for p in people if p.get("name")]
        return InvoiceItem(
            training_id=training.id,
            learner_name=named[0] if len(named) == 1 else "",
            learners=roster_as_text(people),
            learner_roster=dump_roster(people),
            description=training_title_en(training),
            instructor=training.instructor,
            period_start=training.start_date,
            period_end=training.end_date,
            qty=qty,
            unit_price_krw=training.unit_price_krw,
            amount_krw=qty * training.unit_price_krw,
        )

    def filtered_invoices(db, user):
        q = db.query(Invoice)
        if user.role != "admin":
            q = q.filter(Invoice.subsidiary_id == user.subsidiary_id)
        year = request.args.get("year", type=int)
        half = request.args.get("half", type=int)
        status = request.args.get("status", "").strip()
        code = request.args.get("subsidiary", "").strip()
        qtext = request.args.get("q", "").strip()
        company = request.args.get("company", "").strip()
        region = request.args.get("region", "").strip()
        country = request.args.get("country", "").strip()
        if year:
            q = q.filter(Invoice.year == year)
        if half:
            q = q.filter(Invoice.half == half)
        if status:
            q = q.filter(Invoice.status == status)
        admin_entity = user.role == "admin" and bool(code or company or region or country)
        need_sub = bool(qtext or admin_entity)
        if need_sub:
            q = q.join(Subsidiary, Invoice.subsidiary_id == Subsidiary.id)
        if admin_entity:
            if code:
                q = q.filter(Subsidiary.code == code)
            if company:
                q = q.filter(Subsidiary.company == company)
            if region:
                q = q.filter(or_(Subsidiary.region == region, Subsidiary.region_en == region))
            if country:
                q = q.filter(or_(Subsidiary.country == country, Subsidiary.country_en == country))
        if qtext:
            like = f"%{qtext}%"
            clauses = [
                Invoice.invoice_no.ilike(like),
                Invoice.education_name.ilike(like),
            ]
            if need_sub:
                clauses.extend(
                    [
                        Subsidiary.code.ilike(like),
                        Subsidiary.name_en.ilike(like),
                        Subsidiary.name_ko.ilike(like),
                        Subsidiary.company.ilike(like),
                        Subsidiary.country.ilike(like),
                        Subsidiary.country_en.ilike(like),
                        Subsidiary.region.ilike(like),
                        Subsidiary.region_en.ilike(like),
                    ]
                )
            q = q.filter(or_(*clauses))
        return q.order_by(Invoice.created_at.desc())

    def dashboard_stats(db, user, year, half):
        q = db.query(Invoice).filter(Invoice.status != "draft")
        if user.role != "admin":
            q = q.filter(Invoice.subsidiary_id == user.subsidiary_id)
        if year:
            q = q.filter(Invoice.year == year)
        if half:
            q = q.filter(Invoice.half == half)
        invoices = q.all()
        total_krw = sum(i.amount_krw for i in invoices)
        usd = sum(float(i.amount_fx) for i in invoices if i.currency == "USD")
        eur = sum(float(i.amount_fx) for i in invoices if i.currency == "EUR")
        paid_krw = sum(i.amount_krw for i in invoices if i.status == "paid")

        issued = [i for i in invoices if i.status in ("issued", "acknowledged", "paid")]
        by_sub = {}
        for inv in issued:
            by_sub.setdefault(inv.subsidiary_id, []).append(inv)
        ack_entities = 0
        pending_entities = 0
        for _sid, items in by_sub.items():
            if any(i.status == "issued" for i in items):
                pending_entities += 1
            else:
                ack_entities += 1
        entity_total = db.query(Subsidiary).count() if user.role == "admin" else 1
        no_invoice = max(entity_total - len(by_sub), 0)
        counts = {
            "draft": db.query(Invoice).filter(Invoice.status == "draft"),
            "issued": db.query(Invoice).filter(Invoice.status == "issued"),
            "acknowledged": db.query(Invoice).filter(Invoice.status == "acknowledged"),
            "paid": db.query(Invoice).filter(Invoice.status == "paid"),
        }
        if user.role != "admin":
            for k in counts:
                counts[k] = counts[k].filter(Invoice.subsidiary_id == user.subsidiary_id)
        if year:
            for k in counts:
                counts[k] = counts[k].filter(Invoice.year == year)
        if half:
            for k in counts:
                counts[k] = counts[k].filter(Invoice.half == half)
        status_counts = {k: q.count() for k, q in counts.items()}
        recent_q = db.query(Invoice)
        if user.role != "admin":
            recent_q = recent_q.filter(Invoice.subsidiary_id == user.subsidiary_id)
        recent = recent_q.order_by(Invoice.updated_at.desc()).limit(8).all()
        return {
            "total_krw": total_krw,
            "usd": usd,
            "eur": eur,
            "paid_krw": paid_krw,
            "ack_entities": ack_entities,
            "pending_entities": pending_entities,
            "no_invoice": no_invoice,
            "entity_total": entity_total,
            "issued_count": len(issued),
            "status_counts": status_counts,
            "recent": recent,
        }

    # ---------- auth ----------
    @app.route("/lang/<code>")
    def lang(code):
        session["lang"] = "en" if code == "en" else "ko"
        return redirect(request.referrer or url_for("dashboard"))

    @app.route("/login", methods=["GET", "POST"])
    def login():
        if current_user():
            return redirect(url_for("dashboard"))
        if request.method == "POST":
            db = Session()
            username = (request.form.get("username") or "").strip()
            password = request.form.get("password") or ""
            user = db.query(User).filter_by(username=username).first()
            if not user or not check_password_hash(user.password_hash, password):
                flash(t("login_fail"), "danger")
                return render_template("login.html")
            if not user.is_active:
                flash(t("account_disabled"), "danger")
                return render_template("login.html")
            session["user_id"] = user.id
            user.last_login_at = utcnow()
            log_action(db, "user", user.id, "LOGIN", username)
            db.commit()
            if user.must_change_password:
                return redirect(url_for("change_password"))
            return redirect(url_for("dashboard"))
        return render_template("login.html")

    @app.route("/logout")
    def logout():
        session.clear()
        return redirect(url_for("login"))

    @app.route("/change-password", methods=["GET", "POST"])
    @login_required
    def change_password():
        user = current_user()
        if request.method == "POST":
            db = Session()
            user = db.get(User, user.id)
            current = request.form.get("current_password") or ""
            new = request.form.get("new_password") or ""
            confirm = request.form.get("confirm_password") or ""
            if not check_password_hash(user.password_hash, current):
                flash(t("wrong_current"), "danger")
            elif len(new) < 8:
                flash(t("pw_too_short"), "danger")
            elif new != confirm:
                flash(t("pw_mismatch"), "danger")
            elif check_password_hash(user.password_hash, new):
                flash(t("pw_same"), "danger")
            else:
                user.password_hash = generate_password_hash(new)
                user.must_change_password = False
                log_action(db, "user", user.id, "PASSWORD_CHANGE", user.username)
                db.commit()
                flash(t("pw_changed"), "success")
                return redirect(url_for("dashboard"))
        return render_template("change_password.html", forced=user.must_change_password)

    @app.route("/")
    @login_required
    def home():
        return redirect(url_for("dashboard"))

    @app.route("/dashboard")
    @login_required
    def dashboard():
        db = Session()
        user = current_user()
        year = request.args.get("year", type=int) or date.today().year
        if "half" in request.args:
            half = request.args.get("half", type=int)
        else:
            half = half_from_date(date.today())
        stats = dashboard_stats(db, user, year, half)
        years = [y for y in range(date.today().year - 3, date.today().year + 2)]
        return render_template(
            "dashboard.html",
            stats=stats,
            year=year,
            half=half,
            years=years,
            training_count=db.query(Training).count(),
            sub_count=db.query(Subsidiary).count(),
        )

    @app.route("/settings/stamp", methods=["POST"])
    @admin_required
    def stamp_upload():
        db = Session()
        f = request.files.get("stamp")
        if not f or not f.filename or not allowed_stamp(f.filename):
            flash(t("stamp_invalid"), "danger")
            return redirect(url_for("hq_settings"))
        try:
            img = Image.open(f.stream).convert("RGBA")
            img.thumbnail((900, 900))
            STAMP_PATH.parent.mkdir(parents=True, exist_ok=True)
            img.save(STAMP_PATH, "PNG")
        except Exception:
            flash(t("stamp_invalid"), "danger")
            return redirect(url_for("hq_settings"))
        log_action(db, "settings", "stamp", "UPLOAD", secure_filename(f.filename))
        db.commit()
        flash(t("stamp_ok"), "success")
        return redirect(url_for("hq_settings"))

    @app.route("/settings/stamp/delete", methods=["POST"])
    @admin_required
    def stamp_delete():
        db = Session()
        if STAMP_PATH.exists():
            STAMP_PATH.unlink()
        log_action(db, "settings", "stamp", "DELETE", "")
        db.commit()
        flash(t("stamp_removed"), "success")
        return redirect(url_for("hq_settings"))

    @app.route("/settings/stamp/preview")
    @admin_required
    def stamp_preview():
        if not STAMP_PATH.exists():
            return redirect(url_for("hq_settings"))
        return send_file(STAMP_PATH, mimetype="image/png")

    @app.route("/settings", methods=["GET", "POST"])
    @admin_required
    def hq_settings():
        db = Session()
        hq = load_hq_settings()
        if request.method == "POST":
            payload = {
                "HQ_COPY_EMAIL": (request.form.get("hq_copy_email") or "").strip(),
                "PGU_NAME": (request.form.get("pgu_name") or "").strip(),
                "PGU_DEPT": (request.form.get("pgu_dept") or "").strip(),
                "PGU_ADDRESS": (request.form.get("pgu_address") or "").replace("\r", "").strip(),
                "PGU_PHONE": (request.form.get("pgu_phone") or "").strip(),
                "PGU_SIGNATORY": (request.form.get("pgu_signatory") or "").strip(),
                "PGU_TITLE": (request.form.get("pgu_title") or "").strip(),
                "BANK_NAME": (request.form.get("bank_name") or "").strip(),
                "BANK_BRANCH": (request.form.get("bank_branch") or "").strip(),
                "BANK_ACCOUNT_NAME": (request.form.get("bank_account_name") or "").strip(),
                "BANK_ACCOUNT_NO": (request.form.get("bank_account_no") or "").strip(),
                "BANK_SWIFT": (request.form.get("bank_swift") or "").strip(),
                "MAIL_TPL_ISSUE_SUBJECT": (request.form.get("mail_tpl_issue_subject") or "").strip(),
                "MAIL_TPL_ISSUE_BODY": (request.form.get("mail_tpl_issue_body") or "").replace("\r", "").strip(),
                "MAIL_TPL_REMIND_SUBJECT": (request.form.get("mail_tpl_remind_subject") or "").strip(),
                "MAIL_TPL_REMIND_BODY": (request.form.get("mail_tpl_remind_body") or "").replace("\r", "").strip(),
                "MAIL_TPL_RESEND_SUBJECT": (request.form.get("mail_tpl_resend_subject") or "").strip(),
                "MAIL_TPL_RESEND_BODY": (request.form.get("mail_tpl_resend_body") or "").replace("\r", "").strip(),
                "MAIL_TPL_HQ_SUBJECT": (request.form.get("mail_tpl_hq_subject") or "").strip(),
                "MAIL_TPL_HQ_BODY": (request.form.get("mail_tpl_hq_body") or "").replace("\r", "").strip(),
            }
            if not payload["PGU_NAME"] or not payload["PGU_ADDRESS"]:
                flash(t("required"), "danger")
            elif not parse_emails(payload["HQ_COPY_EMAIL"]):
                flash(t("hq_email_invalid"), "danger")
            else:
                hq = save_hq_settings(payload)
                log_action(db, "settings", "hq", "UPDATE", payload["HQ_COPY_EMAIL"])
                db.commit()
                flash(t("saved"), "success")
                return redirect(url_for("hq_settings"))
        return render_template("settings.html", hq=hq, stamp_exists=STAMP_PATH.exists())

    @app.route("/settings/mail/test", methods=["POST"])
    @admin_required
    def mail_test():
        mail = send_test_email(request.form.get("test_to") or "")
        flash_mail_result(mail, "mail_test_ok")
        return redirect(url_for("hq_settings"))

    # ---------- subsidiaries ----------
    def subsidiary_query(db):
        qtext = request.args.get("q", "").strip()
        company = request.args.get("company", "").strip()
        region = request.args.get("region", "").strip()
        country = request.args.get("country", "").strip()
        entity_type = request.args.get("entity_type", "").strip()
        status = request.args.get("status", "").strip()
        q = db.query(Subsidiary)
        if qtext:
            like = f"%{qtext}%"
            q = q.filter(
                or_(
                    Subsidiary.code.ilike(like),
                    Subsidiary.name_ko.ilike(like),
                    Subsidiary.name_en.ilike(like),
                    Subsidiary.country.ilike(like),
                    Subsidiary.region.ilike(like),
                    Subsidiary.region_en.ilike(like),
                    Subsidiary.entity_type.ilike(like),
                    Subsidiary.company.ilike(like),
                )
            )
        if company:
            q = q.filter(Subsidiary.company == company)
        if region:
            q = q.filter(or_(Subsidiary.region == region, Subsidiary.region_en == region))
        if country:
            q = q.filter(or_(Subsidiary.country == country, Subsidiary.country_en == country))
        if entity_type:
            q = q.filter(Subsidiary.entity_type == entity_type)
        if status in ("active", "inactive"):
            q = q.filter(Subsidiary.status == status)
        return q.order_by(Subsidiary.list_no.is_(None), Subsidiary.list_no.asc(), Subsidiary.code.asc(), Subsidiary.id.asc())

    def subsidiary_filter_options(db):
        rows = db.query(Subsidiary).all()
        companies = sorted({(r.company or "").strip() for r in rows if (r.company or "").strip()})
        countries = sorted({(r.country or "").strip() for r in rows if (r.country or "").strip()})
        return {
            "q": request.args.get("q", "").strip(),
            "company": request.args.get("company", "").strip(),
            "region": request.args.get("region", "").strip(),
            "country": request.args.get("country", "").strip(),
            "entity_type": request.args.get("entity_type", "").strip(),
            "status": request.args.get("status", "").strip(),
            "companies": companies,
            "countries": countries,
        }

    @app.route("/subsidiaries")
    @admin_required
    def subsidiaries():
        db = Session()
        rows = subsidiary_query(db).all()
        return render_template(
            "subsidiaries/list.html",
            rows=rows,
            filters=subsidiary_filter_options(db),
            json_ready=MASTER_JSON.exists(),
            regions=REGIONS,
        )

    def save_subsidiary(sub, form, db, is_new=False):
        sub.code = (form.get("code") or sub.code or "").strip().upper()
        try:
            list_no = int(form.get("list_no") or 0)
        except (TypeError, ValueError):
            list_no = 0
        sub.list_no = list_no or None
        sub.company = (form.get("company") or "").strip()
        ko, en = resolve_region(form.get("region"), form.get("region_en"), country=form.get("country"), country_en=form.get("country_en"))
        if not ko:
            raise ValueError("region")
        sub.region = ko
        sub.region_en = en
        sub.country = (form.get("country") or "").strip()
        sub.country_en = (form.get("country_en") or "").strip()
        resolved_type = resolve_entity_type(form.get("entity_type"))
        if resolved_type:
            sub.entity_type = resolved_type
        elif not getattr(sub, "entity_type", None):
            sub.entity_type = ""
        sub.name_ko = (form.get("name_ko") or "").strip()
        sub.name_en = (form.get("name_en") or "").strip()
        sub.address_en = (form.get("address_en") or "").strip()
        sub.notes = (form.get("notes") or "").strip()
        if hasattr(form, "getlist"):
            contact = form.getlist("contact_emails")
            expatriate = form.getlist("hr_emails")
        else:
            contact = []
            expatriate = []
        if contact or (hasattr(form, "getlist") and "contact_emails" in form):
            sub.emails = emails_to_text(contact)
        else:
            emails = (form.get("emails") or "").replace("\r", "")
            sub.emails = emails_to_text(emails_from_text(emails) or [e.strip() for e in emails.replace("\n", ";").split(";") if e.strip()])
        if expatriate or (hasattr(form, "getlist") and "hr_emails" in form):
            sub.hr_emails = emails_to_text(expatriate)
        else:
            sub.hr_emails = emails_to_text(emails_from_text(form.get("hr_emails") or ""))
        sub.phone = (form.get("phone") or "").strip()
        sub.status = form.get("status") or "active"
        same_code = db.query(Subsidiary).filter(Subsidiary.code == sub.code)
        if getattr(sub, "id", None):
            same_code = same_code.filter(Subsidiary.id != sub.id)
        sub.needs_review = bool(same_code.count())
        sub.updated_at = utcnow()
        if is_new:
            db.add(sub)
            db.flush()
            ensure_login_user(db, sub)
        elif sub.user:
            sub.user.is_active = sub.status == "active"

    @app.route("/subsidiaries/new", methods=["GET", "POST"])
    @admin_required
    def subsidiary_new():
        db = Session()
        if request.method == "POST":
            code = (request.form.get("code") or "").strip().upper()
            if not code:
                flash(t("required"), "danger")
            else:
                try:
                    if db.query(Subsidiary).filter_by(code=code).first():
                        flash(t("duplicate_code_review"), "warning")
                    sub = Subsidiary(code=code)
                    save_subsidiary(sub, request.form, db, is_new=True)
                    log_action(db, "subsidiary", f"{sub.id}:{sub.code}", "CREATE", sub.name_en)
                    db.commit()
                    flash(t("saved"), "success")
                    return redirect(url_for("subsidiaries"))
                except Exception:
                    db.rollback()
                    flash(t("required"), "danger")
        return render_template(
            "subsidiaries/form.html",
            row=None,
            regions=REGIONS,
            entity_types=ENTITY_TYPES,
            contact_list=[],
            hr_list=[],
        )

    @app.route("/subsidiaries/<int:sid>/edit", methods=["GET", "POST"])
    @admin_required
    def subsidiary_edit(sid):
        db = Session()
        row = db.get(Subsidiary, sid)
        if not row:
            flash(t("not_found"), "danger")
            return redirect(url_for("subsidiaries"))
        if request.method == "POST":
            try:
                save_subsidiary(row, request.form, db, is_new=False)
                log_action(db, "subsidiary", row.code, "UPDATE", row.name_en)
                db.commit()
                flash(t("saved"), "success")
                return redirect(url_for("subsidiaries"))
            except Exception:
                db.rollback()
                flash(t("required"), "danger")
        return render_template(
            "subsidiaries/form.html",
            row=row,
            regions=REGIONS,
            entity_types=ENTITY_TYPES,
            contact_list=emails_from_text(row.emails),
            hr_list=emails_from_text(row.hr_emails),
        )

    @app.route("/subsidiaries/<int:sid>/reset-password", methods=["POST"])
    @admin_required
    def subsidiary_reset(sid):
        db = Session()
        row = db.get(Subsidiary, sid)
        if row and row.user:
            row.user.password_hash = generate_password_hash(Config.INITIAL_SUBSIDIARY_PASSWORD)
            row.user.must_change_password = True
            log_action(db, "subsidiary", row.code, "PASSWORD_RESET", row.user.username)
            db.commit()
            flash(t("pw_reset_ok"), "success")
        return redirect(url_for("subsidiaries"))

    @app.route("/subsidiaries/<int:sid>/delete", methods=["POST"])
    @admin_required
    def subsidiary_delete(sid):
        db = Session()
        row = db.get(Subsidiary, sid)
        if not row:
            flash(t("not_found"), "danger")
            return redirect(url_for("subsidiaries"))
        linked = db.query(Invoice).filter_by(subsidiary_id=sid).count()
        if linked:
            flash(t("cannot_delete_subsidiary"), "danger")
            return redirect(url_for("subsidiaries"))
        code = row.code
        if row.user:
            db.delete(row.user)
        db.delete(row)
        log_action(db, "subsidiary", code, "DELETE", "")
        db.commit()
        flash(t("subsidiary_deleted"), "success")
        return redirect(url_for("subsidiaries"))

    @app.route("/subsidiaries/export")
    @admin_required
    def subsidiaries_export():
        db = Session()
        rows = subsidiary_query(db).all()
        wb = Workbook()
        ws = wb.active
        ws.title = "subsidiaries"
        headers = [
            "id",
            "list_no",
            "code",
            "company",
            "region",
            "region_en",
            "country",
            "country_en",
            "entity_type",
            "name_ko",
            "name_en",
            "address_en",
            "notes",
            "emails",
            "hr_emails",
            "phone",
            "status",
            "needs_review",
        ]
        ws.append(headers)
        for row in rows:
            ws.append(
                [
                    row.id,
                    getattr(row, "list_no", None) or "",
                    row.code,
                    getattr(row, "company", "") or "",
                    row.region,
                    row.region_en,
                    row.country,
                    row.country_en,
                    getattr(row, "entity_type", "") or "",
                    row.name_ko,
                    row.name_en,
                    row.address_en,
                    getattr(row, "notes", "") or "",
                    row.emails,
                    getattr(row, "hr_emails", "") or "",
                    row.phone,
                    row.status,
                    "Y" if getattr(row, "needs_review", False) else "",
                ]
            )
        bio = BytesIO()
        wb.save(bio)
        bio.seek(0)
        return send_file(
            bio,
            as_attachment=True,
            download_name="subsidiary_master.xlsx",
            mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        )

    @app.route("/subsidiaries/template")
    @admin_required
    def subsidiary_template():
        wb = Workbook()
        ws = wb.active
        ws.title = "subsidiaries"
        headers = [
            "id",
            "list_no",
            "code",
            "company",
            "region",
            "region_en",
            "country",
            "country_en",
            "entity_type",
            "name_ko",
            "name_en",
            "address_en",
            "notes",
            "emails",
            "hr_emails",
            "phone",
            "status",
        ]
        ws.append(headers)
        ws.append(
            [
                "",
                "99",
                "08XX01",
                "포스코",
                "동남아시아",
                "Southeast Asia",
                "샘플",
                "Sample",
                "생산법인",
                "샘플법인",
                "POSCO Sample Co., Ltd.",
                "1 Sample Street, City",
                "",
                "finance@example.com; training@example.com",
                "hr.expat@example.com",
                "+82-32-200-0000",
                "active",
            ]
        )
        bio = BytesIO()
        wb.save(bio)
        bio.seek(0)
        return send_file(
            bio,
            as_attachment=True,
            download_name="subsidiary_master_template.xlsx",
            mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        )

    @app.route("/subsidiaries/upload", methods=["POST"])
    @admin_required
    def subsidiary_upload():
        db = Session()
        f = request.files.get("file")
        if not f or not allowed_file(f.filename):
            flash(t("invalid_file"), "danger")
            return redirect(url_for("subsidiaries"))
        try:
            rows = load_table(f)
        except Exception:
            flash(t("upload_fail"), "danger")
            return redirect(url_for("subsidiaries"))
        created = updated = errors = 0
        for row in rows:
            try:
                sid = str(norm_key(row, "id", "고유번호") or "").strip()
                code = str(norm_key(row, "code", "법인코드")).strip().upper()
                if not code and not sid:
                    errors += 1
                    continue
                sub = None
                if sid.isdigit():
                    sub = db.get(Subsidiary, int(sid))
                if sub is None and code:
                    named = str(norm_key(row, "name_en", "영문법인명", "official_name", "인보이스내공식법인명") or "").strip()
                    cands = db.query(Subsidiary).filter_by(code=code).all()
                    if len(cands) == 1:
                        sub = cands[0]
                    elif named:
                        sub = next((c for c in cands if (c.name_en or "").strip().lower() == named.lower()), None)
                is_new = sub is None
                if is_new:
                    if not code:
                        errors += 1
                        continue
                    sub = Subsidiary(code=code)
                payload = {
                    "code": code or sub.code,
                    "list_no": str(norm_key(row, "list_no", "순번", "no") or getattr(sub, "list_no", "") or ""),
                    "company": str(norm_key(row, "company", "회사") or getattr(sub, "company", "") or ""),
                    "region": str(norm_key(row, "region", "권역") or sub.region or ""),
                    "region_en": str(norm_key(row, "region_en", "권역영문") or getattr(sub, "region_en", "") or ""),
                    "country": str(norm_key(row, "country", "국가") or getattr(sub, "country", "") or ""),
                    "country_en": str(norm_key(row, "country_en", "국가영문") or getattr(sub, "country_en", "") or ""),
                    "entity_type": str(norm_key(row, "entity_type", "해외법인유형", "법인유형") or getattr(sub, "entity_type", "") or ""),
                    "name_ko": str(norm_key(row, "name_ko", "법인명", "정식명") or getattr(sub, "name_ko", "") or ""),
                    "name_en": str(norm_key(row, "name_en", "영문법인명", "official_name", "인보이스내공식법인명") or getattr(sub, "name_en", "") or ""),
                    "address_en": str(norm_key(row, "address_en", "공식영문주소", "법인주소") or getattr(sub, "address_en", "") or ""),
                    "notes": str(norm_key(row, "notes", "기타정보", "비고") or getattr(sub, "notes", "") or ""),
                    "emails": str(norm_key(row, "emails", "이메일", "담당자이메일") or getattr(sub, "emails", "") or ""),
                    "hr_emails": str(norm_key(row, "hr_emails", "인사주재원이메일") or getattr(sub, "hr_emails", "") or ""),
                    "phone": str(norm_key(row, "phone", "연락처") or getattr(sub, "phone", "") or ""),
                    "status": str(norm_key(row, "status", "계정상태") or getattr(sub, "status", "active") or "active"),
                }
                save_subsidiary(sub, payload, db, is_new=is_new)
                if is_new:
                    created += 1
                    log_action(db, "subsidiary", f"{sub.id}:{sub.code}", "BULK_CREATE", payload["name_en"])
                else:
                    updated += 1
                    log_action(db, "subsidiary", f"{sub.id}:{sub.code}", "BULK_UPDATE", payload["name_en"])
            except Exception:
                errors += 1
        mark_duplicate_codes(db)
        db.commit()
        flash(f"{t('upload_ok')} · {t('created_n')} {created} / {t('updated_n')} {updated} / {t('errors')} {errors}", "success")
        return redirect(url_for("subsidiaries"))

    @app.route("/subsidiaries/import-json", methods=["POST"])
    @admin_required
    def subsidiaries_import_json():
        db = Session()
        if not MASTER_JSON.exists():
            flash(t("master_json_missing"), "danger")
            return redirect(url_for("subsidiaries"))
        freeze_issued_parties(db)
        result = import_master_json(db)
        log_action(
            db,
            "subsidiary",
            "master",
            "JSON_IMPORT",
            f"created={result['created']} updated={result['updated']} errors={result['errors']}",
        )
        db.commit()
        flash(
            f"{t('master_json_ok')} · {t('created_n')} {result['created']} / {t('updated_n')} {result['updated']} / {t('errors')} {result['errors']}",
            "success",
        )
        return redirect(url_for("subsidiaries"))

    # ---------- profile (entity) ----------
    @app.route("/profile", methods=["GET", "POST"])
    @login_required
    def profile():
        user = current_user()
        if user.role != "subsidiary" or not user.subsidiary_id:
            flash(t("no_permission"), "danger")
            return redirect(url_for("dashboard"))
        db = Session()
        row = db.get(Subsidiary, user.subsidiary_id)
        if request.method == "POST":
            row.address_en = (request.form.get("address_en") or "").strip()
            emails = (request.form.get("emails") or "").replace("\r", "")
            parsed = parse_emails(emails)
            row.emails = "; ".join(parsed) if parsed else emails.replace("\n", "; ")
            row.phone = (request.form.get("phone") or "").strip()
            row.notes = (request.form.get("notes") or "").strip()
            row.updated_at = utcnow()
            log_action(db, "subsidiary", row.code, "SELF_UPDATE", "address/emails/phone/notes")
            db.commit()
            flash(t("saved"), "success")
            return redirect(url_for("profile"))
        return render_template("profile.html", row=row)

    # ---------- trainings ----------
    @app.route("/trainings")
    @admin_required
    def trainings():
        db = Session()
        rows = db.query(Training).order_by(Training.start_date.desc(), Training.id.desc()).all()
        learner_counts = {row.id: training_headcount(row) for row in rows}
        extras = training_form_extras(db)
        return render_template(
            "trainings/list.html",
            rows=rows,
            learner_counts=learner_counts,
            **extras,
        )

    def save_training(row, form):
        row.title = (form.get("title") or "").strip()
        row.title_en = (form.get("title_en") or "").strip() or row.title
        row.start_date = _as_date(form.get("start_date"))
        row.end_date = _as_date(form.get("end_date"))
        row.instructor = (form.get("instructor") or "").strip()
        row.unit_price_krw = _as_int(form.get("unit_price_krw"))
        row.learners = (form.get("learners") or "").strip()
        raw_count = form.get("learner_count")
        if raw_count in (None, ""):
            grouped, unknown = parse_training_learners(row.learners)
            row.learner_count = sum(len(people) for people in grouped.values()) + len(unknown)
        else:
            row.learner_count = max(_as_int(raw_count, default=0), 0)
        row.updated_at = utcnow()

    def training_headcount(training):
        n = int(getattr(training, "learner_count", 0) or 0)
        if n:
            return n
        grouped, unknown = parse_training_learners(getattr(training, "learners", "") or "")
        return sum(len(people) for people in grouped.values()) + len(unknown)

    def headcount_from_text(text):
        blob = str(text or "").strip()
        if not blob:
            return 0
        found = HEADCOUNT_RE.findall(blob)
        if found:
            return max(int(found[-1]), 1)
        return 1

    def headcount_from_people(people):
        extracted = []
        blobs = []
        for person in people or []:
            blob = " ".join(str(person.get(key) or "") for key in ("name", "email")).strip()
            blobs.append(blob)
            found = HEADCOUNT_RE.findall(blob)
            extracted.append(int(found[-1]) if found else None)
        if not extracted:
            return 0
        if all(n is not None for n in extracted):
            uniq = set(extracted)
            if len(uniq) == 1:
                n = extracted[0]
                if n == 1:
                    return len(extracted)
                if len(set(blobs)) == 1:
                    return n
                return sum(extracted)
            return sum(extracted)
        return len(extracted)

    def entity_headcount_map(training):
        grouped, _ = parse_training_learners(getattr(training, "learners", "") or "")
        return {code: headcount_from_people(people) for code, people in grouped.items()}

    def billed_qty_for_entity(training, entity_code="", people=None, fallback=1):
        code = str(entity_code or "").upper()
        if people is None:
            people = learners_for_entity(training, code) if training and code else []
        n = headcount_from_people(people)
        if n > 0:
            return n
        if training and code:
            grouped, unknown = parse_training_learners(getattr(training, "learners", "") or "")
            if set(grouped.keys()) == {code} and not unknown:
                total = int(getattr(training, "learner_count", 0) or 0)
                if total > 0:
                    return total
        try:
            fb = int(fallback or 0)
        except (TypeError, ValueError):
            fb = 0
        return max(fb, 1)

    def apply_training_headcounts(db, invoice):
        changed = False
        code = entity_code_of(invoice)
        for item in invoice.items or []:
            if not item.training_id:
                continue
            training = db.get(Training, item.training_id)
            if not training:
                continue
            people = load_roster(item) or learners_for_entity(training, code)
            qty = billed_qty_for_entity(training, code, people, fallback=item.qty)
            unit = int(item.unit_price_krw or training.unit_price_krw or 0)
            amount = qty * unit
            if item.qty != qty or item.amount_krw != amount:
                item.qty = qty
                item.unit_price_krw = unit
                item.amount_krw = amount
                changed = True
        if changed:
            apply_invoice_totals(invoice)
            invoice.updated_at = utcnow()
        return changed

    def training_title_en(training):
        return ((getattr(training, "title_en", "") or "") or training.title or "").strip()

    def parse_training_learners(raw):
        grouped = {}
        unknown = []
        for line in str(raw or "").replace("\r", "").splitlines():
            text = line.strip()
            if not text:
                continue
            code, name, email = parse_person(text)
            if not code:
                unknown.append(text)
                continue
            grouped.setdefault(code, [])
            grouped[code].append({"name": name, "email": email})
        return grouped, unknown

    def learners_by_code_map(training):
        grouped, _ = parse_training_learners(getattr(training, "learners", "") or "")
        return {code: roster_as_text(people) for code, people in grouped.items()}

    def learners_for_entity(training, code):
        if not training or not code:
            return []
        grouped, _ = parse_training_learners(training.learners)
        return grouped.get(str(code).upper(), [])

    def entity_code_of(invoice):
        sub = getattr(invoice, "subsidiary", None)
        return (sub.code or "").upper() if sub else ""

    def roster_for_item(item, entity_code=""):
        people = load_roster(item)
        if people:
            return people
        tid = getattr(item, "training_id", None)
        if not tid:
            return []
        sess = object_session(item)
        training = sess.get(Training, tid) if sess else None
        return learners_for_entity(training, entity_code)

    def attach_training_learners(db, inv):
        code = entity_code_of(inv)
        changed = False
        for item in inv.items or []:
            training = db.get(Training, item.training_id) if item.training_id else None
            if training and not load_roster(item):
                people = learners_for_entity(training, code)
                if people:
                    item.learners = roster_as_text(people)
                    item.learner_roster = dump_roster(people)
                    names = [p["name"] for p in people if p.get("name")]
                    item.learner_name = names[0] if len(names) == 1 else ""
                    changed = True
        if apply_training_headcounts(db, inv):
            changed = True
        return changed

    def invoice_roster(invoice):
        code = entity_code_of(invoice)
        rows = []
        for item in invoice.items or []:
            program = (item.description or invoice.education_name or "").strip()
            for person in roster_for_item(item, code):
                rows.append(
                    {
                        "program": program,
                        "name": person.get("name") or "",
                        "email": person.get("email") or "",
                    }
                )
        return rows

    def latest_fx(db, training=None):
        q = db.query(Invoice)
        if training and getattr(training, "id", None):
            row = (
                q.filter(Invoice.training_id == training.id)
                .order_by(Invoice.id.desc())
                .first()
            )
            if row:
                return row.currency, row.exchange_rate, row.rate_date
        row = q.order_by(Invoice.id.desc()).first()
        if row:
            return row.currency, row.exchange_rate, row.rate_date
        return "USD", "", date.today()

    def training_form_extras(db, row=None):
        currency, rate, rate_date = latest_fx(db, row)
        return {
            "fx_currency": currency or "USD",
            "fx_rate": f"{float(rate):.2f}" if rate not in (None, "") else "",
            "fx_rate_date": (rate_date or date.today()).strftime("%Y-%m-%d") if hasattr(rate_date, "strftime") else str(rate_date or date.today())[:10],
            "learner_count_value": training_headcount(row) if row else "",
        }

    def parse_bulk_fx(form):
        try:
            rate = _as_rate(form.get("exchange_rate") or 0)
        except (TypeError, ValueError):
            rate = 0
        currency = (form.get("currency") or "USD").upper()
        if currency not in ("USD", "EUR"):
            currency = "USD"
        rate_date = _as_date(form.get("rate_date") or date.today())
        return currency, rate, rate_date

    def issued_training_ids(db, subsidiary_id, training_ids):
        if not training_ids:
            return set()
        item_rows = (
            db.query(InvoiceItem.training_id)
            .join(Invoice)
            .filter(
                Invoice.subsidiary_id == subsidiary_id,
                Invoice.status != "draft",
                InvoiceItem.training_id.in_(training_ids),
            )
            .all()
        )
        header_rows = (
            db.query(Invoice.training_id)
            .filter(
                Invoice.subsidiary_id == subsidiary_id,
                Invoice.status != "draft",
                Invoice.training_id.in_(training_ids),
            )
            .all()
        )
        out = {row[0] for row in item_rows if row[0]}
        out.update(row[0] for row in header_rows if row[0])
        return out

    def replaceable_drafts(db, subsidiary_id, year, half, training_ids):
        drafts = (
            db.query(Invoice)
            .filter(
                Invoice.subsidiary_id == subsidiary_id,
                Invoice.year == year,
                Invoice.half == half,
                Invoice.status == "draft",
            )
            .order_by(Invoice.id.asc())
            .all()
        )
        selected = set(training_ids)
        found = []
        for inv in drafts:
            item_tids = {it.training_id for it in (inv.items or []) if it.training_id}
            if inv.training_id:
                item_tids.add(inv.training_id)
            if item_tids and item_tids <= selected:
                found.append(inv)
        return found

    def bulk_invoices_from_trainings(db, trainings, form):
        currency, rate, rate_date = parse_bulk_fx(form)
        if rate <= 0:
            return None
        buckets = {}
        unknown = []
        subs_by_code = {(row.code or "").upper(): row for row in db.query(Subsidiary).all()}
        for training in trainings:
            grouped, bad = parse_training_learners(training.learners)
            unknown.extend(bad)
            if not grouped:
                continue
            year = training.start_date.year
            half = half_from_date(training.start_date)
            for code, people in grouped.items():
                sub = subs_by_code.get(code)
                if not sub:
                    unknown.append(code)
                    continue
                buckets.setdefault((sub.id, year, half), []).append((training, people))
        if not buckets:
            return 0, 0, 0, list(dict.fromkeys(unknown)), True
        created = updated = skipped = 0
        for (sub_id, year, half), lines in buckets.items():
            tid_list = list(dict.fromkeys(training.id for training, _ in lines))
            blocked = issued_training_ids(db, sub_id, tid_list)
            pending = [(training, people) for training, people in lines if training.id not in blocked]
            if not pending:
                skipped += 1
                continue
            pending_ids = [training.id for training, _ in pending]
            drafts = replaceable_drafts(db, sub_id, year, half, pending_ids)
            if drafts:
                inv = drafts[0]
                for extra in drafts[1:]:
                    db.delete(extra)
                inv.items.clear()
                db.flush()
                updated += 1
            else:
                inv_no, seq = next_invoice_no(db, year, half)
                first = pending[0][0]
                inv = Invoice(
                    invoice_no=inv_no,
                    year=year,
                    half=half,
                    seq=seq,
                    subsidiary_id=sub_id,
                    training_id=None,
                    education_name=training_title_en(first),
                    period_start=first.start_date,
                    period_end=first.end_date,
                    instructor=first.instructor,
                    qty=1,
                    unit_price_krw=first.unit_price_krw,
                    amount_krw=0,
                    currency=currency,
                    exchange_rate=rate,
                    rate_date=rate_date,
                    amount_fx=0,
                    vat_rate=0.0,
                    status="draft",
                    document_date=date.today(),
                )
                db.add(inv)
                db.flush()
                created += 1
            inv.currency = currency
            inv.exchange_rate = rate
            inv.rate_date = rate_date
            inv.document_date = date.today()
            inv.year = year
            inv.half = half
            inv.updated_at = utcnow()
            for training, people in pending:
                inv.items.append(invoice_item_from_training(training, people))
            apply_invoice_header_from_items(inv)
            sub = db.get(Subsidiary, sub_id)
            log_action(
                db,
                "invoice",
                inv.invoice_no,
                "BULK_FROM_TRAININGS",
                f"{sub.code if sub else sub_id} programs={len(pending)}",
            )
        return created, updated, skipped, list(dict.fromkeys(unknown)), False

    def flash_auto_invoices(result):
        if result is None:
            flash(t("auto_invoice_need_rate"), "danger")
            return
        created, updated, skipped, unknown, empty = result
        if empty and not created and not updated:
            flash(t("bulk_no_learners"), "warning")
            if unknown:
                flash(f"{t('auto_invoice_unknown')}: {', '.join(unknown)}", "warning")
            return
        bits = []
        if created:
            bits.append(f"{t('auto_invoice_created')} {created}")
        if updated:
            bits.append(f"{t('auto_invoice_updated')} {updated}")
        if bits:
            flash(" · ".join(bits), "success")
        elif not skipped:
            flash(t("bulk_none"), "warning")
        if skipped:
            flash(t("auto_invoice_skipped"), "warning")
        if unknown:
            flash(f"{t('auto_invoice_unknown')}: {', '.join(unknown)}", "warning")

    @app.route("/trainings/new", methods=["GET", "POST"])
    @admin_required
    def training_new():
        db = Session()
        if request.method == "POST":
            try:
                row = Training(
                    title="",
                    title_en="",
                    start_date=date.today(),
                    end_date=date.today(),
                    instructor="",
                    unit_price_krw=0,
                )
                save_training(row, request.form)
                db.add(row)
                db.flush()
                log_action(db, "training", row.title, "CREATE", row.instructor)
                db.commit()
                flash(t("saved"), "success")
                return redirect(url_for("trainings"))
            except Exception:
                db.rollback()
                flash(t("required"), "danger")
        extras = training_form_extras(db)
        return render_template("trainings/form.html", row=None, **extras)

    @app.route("/trainings/<int:tid>/edit", methods=["GET", "POST"])
    @admin_required
    def training_edit(tid):
        db = Session()
        row = db.get(Training, tid)
        if not row:
            return redirect(url_for("trainings"))
        if request.method == "POST":
            try:
                save_training(row, request.form)
                log_action(db, "training", row.id, "UPDATE", row.title)
                db.commit()
                flash(t("saved"), "success")
                return redirect(url_for("trainings"))
            except Exception:
                db.rollback()
                flash(t("required"), "danger")
        extras = training_form_extras(db, row)
        return render_template("trainings/form.html", row=row, **extras)

    @app.route("/trainings/<int:tid>/delete", methods=["POST"])
    @admin_required
    def training_delete(tid):
        db = Session()
        row = db.get(Training, tid)
        if row:
            linked = (
                db.query(Invoice).filter_by(training_id=tid).count()
                or db.query(InvoiceItem).filter_by(training_id=tid).count()
            )
            if linked:
                flash(t("cannot_delete"), "danger")
            else:
                db.delete(row)
                db.commit()
                flash(t("saved"), "success")
        return redirect(url_for("trainings"))

    @app.route("/trainings/bulk-invoices", methods=["POST"])
    @admin_required
    def training_bulk_invoices():
        db = Session()
        ids = []
        for value in request.form.getlist("training_ids"):
            if str(value).isdigit():
                ids.append(int(value))
        if not ids:
            flash(t("bulk_need_training"), "danger")
            return redirect(url_for("trainings"))
        rows = (
            db.query(Training)
            .filter(Training.id.in_(ids))
            .order_by(Training.start_date.asc(), Training.id.asc())
            .all()
        )
        try:
            result = bulk_invoices_from_trainings(db, rows, request.form)
            if result is None:
                db.rollback()
                flash_auto_invoices(result)
                return redirect(url_for("trainings"))
            db.commit()
            flash_auto_invoices(result)
            created, updated, *_rest = result
            if created or updated:
                return redirect(url_for("invoices"))
        except Exception:
            db.rollback()
            flash(t("required"), "danger")
        return redirect(url_for("trainings"))

    @app.route("/api/trainings/<int:tid>")
    @admin_required
    def api_training(tid):
        db = Session()
        row = db.get(Training, tid)
        if not row:
            return {"ok": False}, 404
        return {
            "ok": True,
            "title": row.title,
            "title_en": getattr(row, "title_en", "") or row.title,
            "start_date": row.start_date.strftime("%Y-%m-%d"),
            "end_date": row.end_date.strftime("%Y-%m-%d"),
            "instructor": row.instructor,
            "unit_price_krw": row.unit_price_krw,
        }

    # ---------- invoices ----------
    @app.route("/invoices")
    @login_required
    def invoices():
        db = Session()
        user = current_user()
        q = filtered_invoices(db, user)
        rows = q.all()
        years = sorted({i.year for i in db.query(Invoice.year).all()} | {date.today().year})
        subs = db.query(Subsidiary).order_by(Subsidiary.code).all() if user.role == "admin" else []
        extras = training_form_extras(db) if user.role == "admin" else {}
        entity_opts = subsidiary_filter_options(db) if user.role == "admin" else {}
        return render_template(
            "invoices/list.html",
            rows=rows,
            years=years,
            subs=subs,
            regions=REGIONS,
            filters={
                "year": request.args.get("year", ""),
                "half": request.args.get("half", ""),
                "status": request.args.get("status", ""),
                "subsidiary": request.args.get("subsidiary", ""),
                "q": request.args.get("q", ""),
                "company": entity_opts.get("company", ""),
                "region": entity_opts.get("region", ""),
                "country": entity_opts.get("country", ""),
                "companies": entity_opts.get("companies", []),
                "countries": entity_opts.get("countries", []),
            },
            **extras,
        )

    @app.route("/invoices/export")
    @login_required
    def invoices_export():
        db = Session()
        user = current_user()
        rows = filtered_invoices(db, user).all()
        buf = StringIO()
        w = csv.writer(buf)
        w.writerow(
            [
                "invoice_no",
                "year",
                "half",
                "status",
                "entity_code",
                "entity_name_en",
                "education_name",
                "period_start",
                "period_end",
                "instructor",
                "qty",
                "unit_price_krw",
                "amount_krw",
                "currency",
                "exchange_rate",
                "rate_date",
                "amount_fx",
                "vat_rate",
                "document_date",
                "issued_at",
                "acknowledged_at",
                "acknowledged_by",
                "payment_date",
                "paid_memo",
                "updated_at",
            ]
        )
        for i in rows:
            w.writerow(
                [
                    i.invoice_no,
                    i.year,
                    f"{i.half:02d}",
                    i.status,
                    i.subsidiary.code,
                    i.subsidiary.name_en,
                    i.education_name,
                    i.period_start.strftime("%Y-%m-%d"),
                    i.period_end.strftime("%Y-%m-%d"),
                    i.instructor,
                    i.qty,
                    i.unit_price_krw,
                    i.amount_krw,
                    i.currency,
                    f"{float(i.exchange_rate):.2f}",
                    i.rate_date.strftime("%Y-%m-%d"),
                    f"{float(i.amount_fx):.2f}",
                    i.vat_rate,
                    i.document_date.strftime("%Y-%m-%d"),
                    i.issued_at.strftime("%Y-%m-%d %H:%M") if i.issued_at else "",
                    i.acknowledged_at.strftime("%Y-%m-%d %H:%M") if i.acknowledged_at else "",
                    i.acknowledged_by or "",
                    i.payment_date.strftime("%Y-%m-%d") if i.payment_date else "",
                    i.paid_memo or "",
                    i.updated_at.strftime("%Y-%m-%d %H:%M") if i.updated_at else "",
                ]
            )
        data = ("\ufeff" + buf.getvalue()).encode("utf-8")
        return Response(
            data,
            mimetype="text/csv",
            headers={"Content-Disposition": "attachment; filename=pgu_invoices.csv"},
        )

    def fill_invoice_from_form(db, inv, form, is_new=False):
        sub_id = form.get("subsidiary_id", type=int)
        sub = db.get(Subsidiary, sub_id)
        if not sub:
            raise ValueError("subsidiary")
        inv.subsidiary_id = sub.id
        inv.currency = (form.get("currency") or "USD").upper()
        if inv.currency not in ("USD", "EUR"):
            inv.currency = "USD"
        inv.exchange_rate = _as_rate(form.get("exchange_rate") or 0)
        inv.rate_date = _as_date(form.get("rate_date") or date.today())
        inv.document_date = _as_date(form.get("document_date") or date.today())
        inv.vat_rate = 0.0

        titles = form.getlist("item_education_name")
        instructors = form.getlist("item_instructor")
        starts = form.getlist("item_period_start")
        ends = form.getlist("item_period_end")
        qtys = form.getlist("item_qty")
        units = form.getlist("item_unit_price_krw")
        learners_list = form.getlist("item_learners")
        training_ids = form.getlist("item_training_id")

        programs = []
        for i, raw_title in enumerate(titles):
            title = (raw_title or "").strip()
            if not title:
                continue
            qty = _as_int(qtys[i] if i < len(qtys) else 1, default=1) or 1
            unit = _as_int(units[i] if i < len(units) else 0)
            raw_learners = learners_list[i] if i < len(learners_list) else ""
            people = parse_people_lines(raw_learners)
            tid_raw = training_ids[i] if i < len(training_ids) else ""
            tid = int(tid_raw) if str(tid_raw).isdigit() else None
            training = db.get(Training, tid) if tid else None
            if not people and training:
                people = learners_for_entity(training, sub.code)
            if training:
                qty = billed_qty_for_entity(training, sub.code, people, fallback=qty)
            names = [p["name"] for p in people if p.get("name")]
            ps = _as_date(starts[i] if i < len(starts) and starts[i] else date.today())
            pe = _as_date(ends[i] if i < len(ends) and ends[i] else ps)
            instructor = (instructors[i] if i < len(instructors) else "") or ""
            programs.append(
                {
                    "title": title,
                    "instructor": instructor.strip(),
                    "start": ps,
                    "end": pe,
                    "qty": qty,
                    "unit": unit,
                    "learners": roster_as_text(people),
                    "learner_name": names[0] if len(names) == 1 else "",
                    "learner_roster": dump_roster(people),
                    "training_id": tid,
                }
            )
        if not programs:
            raise ValueError("programs")

        if not is_new:
            inv.items.clear()
            db.flush()
        for prog in programs:
            inv.items.append(
                InvoiceItem(
                    training_id=prog["training_id"],
                    learner_name=prog["learner_name"],
                    learners=prog["learners"],
                    learner_roster=prog.get("learner_roster") or "",
                    description=prog["title"],
                    instructor=prog["instructor"],
                    period_start=prog["start"],
                    period_end=prog["end"],
                    qty=prog["qty"],
                    unit_price_krw=prog["unit"],
                    amount_krw=prog["qty"] * prog["unit"],
                )
            )

        titles_clean = [p["title"] for p in programs]
        if len(titles_clean) == 1:
            inv.education_name = titles_clean[0]
        else:
            joined = " / ".join(titles_clean)
            inv.education_name = joined if len(joined) <= 300 else f"{titles_clean[0]} (+{len(titles_clean) - 1})"
        inv.period_start = min(p["start"] for p in programs)
        inv.period_end = max(p["end"] for p in programs)
        inv.instructor = ", ".join(dict.fromkeys(p["instructor"] for p in programs if p["instructor"]))
        unique_tids = list(dict.fromkeys(p["training_id"] for p in programs if p["training_id"]))
        inv.training_id = unique_tids[0] if len(unique_tids) == 1 else None
        inv.year = inv.period_start.year
        inv.half = half_from_date(inv.period_start)
        apply_invoice_totals(inv)
        if is_new:
            inv.invoice_no, inv.seq = next_invoice_no(db, inv.year, inv.half)
            inv.status = "draft"
            db.add(inv)
        inv.updated_at = utcnow()

    def program_rows_from_invoice(row):
        items = list(row.items or [])
        if not items:
            return [
                {
                    "training_id": row.training_id or "",
                    "education_name": row.education_name or "",
                    "instructor": row.instructor or "",
                    "period_start": row.period_start.strftime("%Y-%m-%d") if row.period_start else "",
                    "period_end": row.period_end.strftime("%Y-%m-%d") if row.period_end else "",
                    "qty": row.qty or 1,
                    "unit_price_krw": row.unit_price_krw or 0,
                    "learners": "",
                }
            ]
        if any(getattr(it, "period_start", None) for it in items):
            rows = []
            for it in items:
                ps = it.period_start or row.period_start
                pe = it.period_end or row.period_end
                rows.append(
                    {
                        "training_id": it.training_id or "",
                        "education_name": it.description or row.education_name,
                        "instructor": it.instructor or row.instructor or "",
                        "period_start": ps.strftime("%Y-%m-%d") if ps else "",
                        "period_end": pe.strftime("%Y-%m-%d") if pe else "",
                        "qty": it.qty or 1,
                        "unit_price_krw": it.unit_price_krw or 0,
                        "learners": roster_as_text(roster_for_item(it, entity_code_of(row))),
                    }
                )
            return rows
        from collections import OrderedDict

        groups = OrderedDict()
        for it in items:
            key = (it.description or row.education_name, it.unit_price_krw)
            g = groups.setdefault(
                key,
                {"qty": 0, "people": [], "unit": it.unit_price_krw, "desc": it.description or row.education_name},
            )
            g["qty"] += it.qty or 1
            g["people"].extend(roster_for_item(it, entity_code_of(row)))
        return [
            {
                "training_id": row.training_id or "",
                "education_name": g["desc"],
                "instructor": row.instructor or "",
                "period_start": row.period_start.strftime("%Y-%m-%d") if row.period_start else "",
                "period_end": row.period_end.strftime("%Y-%m-%d") if row.period_end else "",
                "qty": g["qty"],
                "unit_price_krw": g["unit"],
                "learners": roster_as_text(g["people"]),
            }
            for g in groups.values()
        ]

    def trainings_payload(trains):
        rows = []
        for tr in trains:
            by_code = learners_by_code_map(tr)
            counts_by_code = entity_headcount_map(tr)
            rows.append(
                {
                    "id": tr.id,
                    "training_id": tr.id,
                    "education_name": training_title_en(tr),
                    "education_name_ko": tr.title,
                    "instructor": tr.instructor,
                    "period_start": tr.start_date.strftime("%Y-%m-%d"),
                    "period_end": tr.end_date.strftime("%Y-%m-%d"),
                    "qty": 1,
                    "learner_count": int(getattr(tr, "learner_count", 0) or 0),
                    "counts_by_code": counts_by_code,
                    "unit_price_krw": tr.unit_price_krw,
                    "learners": "",
                    "learners_by_code": by_code,
                }
            )
        return rows

    @app.route("/invoices/new", methods=["GET", "POST"])
    @admin_required
    def invoice_new():
        db = Session()
        if request.method == "POST":
            try:
                inv = Invoice(
                    invoice_no="TEMP",
                    year=date.today().year,
                    half=1,
                    seq=0,
                    subsidiary_id=0,
                    education_name="",
                    period_start=date.today(),
                    period_end=date.today(),
                    qty=1,
                    unit_price_krw=0,
                    amount_krw=0,
                    currency="USD",
                    exchange_rate=1,
                    rate_date=date.today(),
                    amount_fx=0,
                    document_date=date.today(),
                )
                fill_invoice_from_form(db, inv, request.form, is_new=True)
                log_action(db, "invoice", inv.invoice_no, "CREATE", inv.education_name)
                db.commit()
                flash(t("saved"), "success")
                return redirect(url_for("invoice_detail", iid=inv.id))
            except Exception:
                db.rollback()
                flash(t("required"), "danger")
        subs = db.query(Subsidiary).filter_by(status="active").order_by(Subsidiary.code).all()
        trains = db.query(Training).order_by(Training.start_date.desc()).all()
        extras = training_form_extras(db)
        entity_opts = subsidiary_filter_options(db)
        return render_template(
            "invoices/form.html",
            row=None,
            subs=subs,
            trains=trains,
            catalog=trainings_payload(trains),
            program_rows=[],
            default_currency=extras["fx_currency"],
            default_rate=extras["fx_rate"],
            default_rate_date=extras["fx_rate_date"],
            regions=REGIONS,
            entity_companies=entity_opts["companies"],
            entity_countries=entity_opts["countries"],
        )

    @app.route("/invoices/<int:iid>/edit", methods=["GET", "POST"])
    @admin_required
    def invoice_edit(iid):
        db = Session()
        row = db.get(Invoice, iid)
        if not row:
            flash(t("not_found"), "danger")
            return redirect(url_for("invoices"))
        if request.method == "POST":
            try:
                fill_invoice_from_form(db, row, request.form, is_new=False)
                if row.status != "draft":
                    rebuild_pdf(row)
                log_action(db, "invoice", row.invoice_no, "UPDATE", f"status={row.status}")
                db.commit()
                flash(t("saved"), "success")
                return redirect(url_for("invoice_detail", iid=row.id))
            except Exception:
                db.rollback()
                flash(t("required"), "danger")
        subs = db.query(Subsidiary).order_by(Subsidiary.code).all()
        trains = db.query(Training).order_by(Training.start_date.desc()).all()
        if attach_training_learners(db, row):
            db.commit()
        entity_opts = subsidiary_filter_options(db)
        return render_template(
            "invoices/form.html",
            row=row,
            subs=subs,
            trains=trains,
            catalog=trainings_payload(trains),
            program_rows=program_rows_from_invoice(row),
            default_rate_date=date.today().strftime("%Y-%m-%d"),
            regions=REGIONS,
            entity_companies=entity_opts["companies"],
            entity_countries=entity_opts["countries"],
        )

    @app.route("/invoices/<int:iid>")
    @login_required
    def invoice_detail(iid):
        db = Session()
        user = current_user()
        row = db.get(Invoice, iid)
        if not row:
            flash(t("not_found"), "danger")
            return redirect(url_for("invoices"))
        if user.role != "admin" and row.subsidiary_id != user.subsidiary_id:
            flash(t("no_permission"), "danger")
            return redirect(url_for("invoices"))
        if attach_training_learners(db, row):
            db.commit()
        logs = (
            db.query(AuditLog)
            .filter(AuditLog.entity_type == "invoice", AuditLog.entity_id == row.invoice_no)
            .order_by(AuditLog.created_at.desc())
            .all()
        )
        return render_template(
            "invoices/detail.html",
            row=row,
            logs=logs,
            roster=invoice_roster(row),
        )

    def parse_selected_invoice_ids(form):
        ids = []
        seen = set()
        for raw in form.getlist("invoice_ids"):
            try:
                iid = int(raw)
            except (TypeError, ValueError):
                continue
            if iid in seen:
                continue
            seen.add(iid)
            ids.append(iid)
        return ids

    def invoice_pdf_files(row):
        paths = []
        if row.pdf_path:
            paths.append(Path(row.pdf_path))
        paths.append(invoice_pdf_path(row))
        seen = set()
        unique = []
        for path in paths:
            key = str(path.resolve()) if path else ""
            if not key or key in seen:
                continue
            seen.add(key)
            unique.append(path)
        return unique

    def unlink_paths(paths):
        for path in paths:
            try:
                if path.is_file():
                    path.unlink()
            except OSError:
                pass

    def delete_invoice_row(db, row):
        paths = invoice_pdf_files(row)
        log_action(db, "invoice", row.invoice_no, "DELETE", row.education_name)
        db.delete(row)
        return paths

    @app.route("/invoices/<int:iid>/delete", methods=["POST"])
    @admin_required
    def invoice_delete(iid):
        db = Session()
        row = db.get(Invoice, iid)
        if not row:
            flash(t("not_found"), "danger")
            return redirect(url_for("invoices"))
        paths = delete_invoice_row(db, row)
        db.commit()
        unlink_paths(paths)
        flash(t("invoice_deleted"), "success")
        return redirect(url_for("invoices"))

    @app.route("/invoices/<int:iid>/compose/<kind>", methods=["GET"])
    @admin_required
    def invoice_compose(iid, kind):
        db = Session()
        row = db.get(Invoice, iid)
        if not row:
            return redirect(url_for("invoices"))
        if kind == "issue":
            if row.status != "draft":
                flash(t("issue_disabled"), "warning")
                return redirect(url_for("invoice_detail", iid=iid))
            apply_training_headcounts(db, row)
            db.commit()
            action = url_for("invoice_issue", iid=iid)
        elif kind == "remind":
            if row.status not in ("issued", "acknowledged"):
                flash(t("issued_only_remind"), "warning")
                return redirect(url_for("invoice_detail", iid=iid))
            action = url_for("invoice_remind", iid=iid)
        elif kind == "resend":
            if row.status == "draft":
                return redirect(url_for("invoice_detail", iid=iid))
            action = url_for("invoice_resend", iid=iid)
        else:
            flash(t("not_found"), "danger")
            return redirect(url_for("invoice_detail", iid=iid))
        if kind != "hq" and not contact_emails(row.subsidiary):
            flash(t("no_email"), "danger")
            return redirect(url_for("invoice_detail", iid=iid))
        return render_mail_compose(
            kind,
            [row],
            action,
            url_for("invoice_detail", iid=iid),
            fill=True,
        )

    @app.route("/invoices/<int:iid>/issue", methods=["POST"])
    @admin_required
    def invoice_issue(iid):
        db = Session()
        row = db.get(Invoice, iid)
        if not row:
            return redirect(url_for("invoices"))
        if row.status != "draft":
            flash(t("issue_disabled"), "warning")
            return redirect(url_for("invoice_detail", iid=iid))
        to_list, cc_list = parse_compose_recipients(request.form, row, "issue")
        if not to_list:
            flash(t("no_email"), "danger")
            return redirect(url_for("invoice_detail", iid=iid))
        if request.form.get("compose_ready") != "1":
            return redirect(url_for("invoice_compose", iid=iid, kind="issue"))
        apply_training_headcounts(db, row)
        snapshot_invoice_party(row)
        path = rebuild_pdf(row)
        mail = send_invoice_email(
            row,
            path,
            kind="issue",
            subject=request.form.get("mail_subject"),
            body_text=request.form.get("mail_body"),
            to_recipients=to_list,
            cc_recipients=cc_list,
        )
        if not flash_mail_result(mail, "issued_ok", commit=False):
            return redirect(url_for("invoice_detail", iid=iid))
        row.status = "issued"
        row.issued_at = utcnow()
        row.issued_by = current_user().username
        row.document_date = date.today()
        log_action(
            db,
            "invoice",
            row.invoice_no,
            "ISSUE",
            mail_log_detail(mail),
        )
        db.commit()
        return redirect(url_for("invoice_detail", iid=iid))

    @app.route("/invoices/<int:iid>/acknowledge", methods=["POST"])
    @login_required
    def invoice_ack(iid):
        db = Session()
        user = current_user()
        row = db.get(Invoice, iid)
        if not row:
            return redirect(url_for("invoices"))
        if user.role != "subsidiary" or row.subsidiary_id != user.subsidiary_id:
            flash(t("no_permission"), "danger")
            return redirect(url_for("invoice_detail", iid=iid))
        if row.status != "issued":
            return redirect(url_for("invoice_detail", iid=iid))
        row.status = "acknowledged"
        row.acknowledged_at = utcnow()
        row.acknowledged_by = user.username
        log_action(db, "invoice", row.invoice_no, "ACKNOWLEDGE", user.username)
        db.commit()
        flash(t("acked_ok"), "success")
        return redirect(url_for("invoice_detail", iid=iid))

    @app.route("/invoices/<int:iid>/paid", methods=["POST"])
    @admin_required
    def invoice_paid(iid):
        db = Session()
        row = db.get(Invoice, iid)
        if not row:
            return redirect(url_for("invoices"))
        if row.status == "draft":
            flash(t("issue_disabled"), "warning")
            return redirect(url_for("invoice_detail", iid=iid))
        if row.status == "paid":
            flash(t("already_paid"), "info")
            return redirect(url_for("invoice_detail", iid=iid))
        row.status = "paid"
        row.paid_at = utcnow()
        row.paid_by = current_user().username
        row.payment_date = _as_date(request.form.get("payment_date") or date.today())
        row.paid_memo = (request.form.get("paid_memo") or "").strip()
        log_action(db, "invoice", row.invoice_no, "PAID", f"{row.payment_date} {row.paid_memo}")
        db.commit()
        flash(t("paid_ok"), "success")
        return redirect(url_for("invoice_detail", iid=iid))

    @app.route("/invoices/<int:iid>/remind", methods=["POST"])
    @admin_required
    def invoice_remind(iid):
        db = Session()
        row = db.get(Invoice, iid)
        if not row or row.status not in ("issued", "acknowledged"):
            flash(t("issued_only_remind"), "warning")
            return redirect(url_for("invoice_detail", iid=iid) if row else url_for("invoices"))
        to_list, cc_list = parse_compose_recipients(request.form, row, "remind")
        if not to_list:
            flash(t("no_email"), "danger")
            return redirect(url_for("invoice_detail", iid=iid))
        if request.form.get("compose_ready") != "1":
            return redirect(url_for("invoice_compose", iid=iid, kind="remind"))
        path = rebuild_pdf(row)
        mail = send_invoice_email(
            row,
            path,
            kind="remind",
            subject=request.form.get("mail_subject"),
            body_text=request.form.get("mail_body"),
            to_recipients=to_list,
            cc_recipients=cc_list,
        )
        if flash_mail_result(mail, "remind_ok", commit=False):
            log_action(db, "invoice", row.invoice_no, "REMIND", mail_log_detail(mail))
            db.commit()
        return redirect(url_for("invoice_detail", iid=iid))

    @app.route("/invoices/<int:iid>/resend", methods=["POST"])
    @admin_required
    def invoice_resend(iid):
        db = Session()
        row = db.get(Invoice, iid)
        if not row or row.status == "draft":
            return redirect(url_for("invoices"))
        to_list, cc_list = parse_compose_recipients(request.form, row, "resend")
        if not to_list:
            flash(t("no_email"), "danger")
            return redirect(url_for("invoice_detail", iid=iid))
        if request.form.get("compose_ready") != "1":
            return redirect(url_for("invoice_compose", iid=iid, kind="resend"))
        path = rebuild_pdf(row)
        mail = send_invoice_email(
            row,
            path,
            kind="resend",
            subject=request.form.get("mail_subject"),
            body_text=request.form.get("mail_body"),
            to_recipients=to_list,
            cc_recipients=cc_list,
        )
        if flash_mail_result(mail, "mail_resent", commit=False):
            log_action(db, "invoice", row.invoice_no, "RESEND", mail_log_detail(mail))
            db.commit()
        return redirect(url_for("invoice_detail", iid=iid))

    @app.route("/invoices/bulk-remind", methods=["POST"])
    @admin_required
    def invoices_bulk_remind():
        db = Session()
        ids = parse_selected_invoice_ids(request.form)
        if not ids:
            flash(t("bulk_remind_none"), "warning")
            return redirect(url_for("invoices", **request.args))
        rows = (
            db.query(Invoice)
            .filter(Invoice.id.in_(ids), Invoice.status == "issued")
            .order_by(Invoice.year.asc(), Invoice.half.asc(), Invoice.seq.asc())
            .all()
        )
        if not rows:
            flash(t("bulk_remind_none"), "warning")
            return redirect(url_for("invoices", **request.args))
        if request.form.get("compose_ready") != "1":
            return render_mail_compose(
                "remind",
                rows,
                url_for("invoices_bulk_remind", **request.args),
                url_for("invoices", **request.args),
                fill=False,
            )
        subject = request.form.get("mail_subject")
        body_text = request.form.get("mail_body")
        sent = failed = skipped = 0
        via = "outbox"
        for row in rows:
            to_list, cc_list = parse_compose_recipients(request.form, row, "remind")
            if not to_list:
                skipped += 1
                continue
            try:
                path = rebuild_pdf(row)
                mail = send_invoice_email(
                    row,
                    path,
                    kind="remind",
                    subject=subject,
                    body_text=body_text,
                    to_recipients=to_list,
                    cc_recipients=cc_list,
                )
                if mail.get("ok"):
                    sent += 1
                    if mail.get("via") == "smtp":
                        via = "smtp"
                    log_action(db, "invoice", row.invoice_no, "REMIND", mail_log_detail(mail))
                else:
                    failed += 1
            except Exception:
                failed += 1
        db.commit()
        if sent:
            flash(f"{t('bulk_remind_ok')} {sent}{t('rows')}", "success")
            if via == "outbox":
                flash(t("email_saved_outbox"), "info")
        if skipped:
            flash(f"{t('no_email')} {skipped}{t('rows')}", "warning")
        if failed:
            flash(f"{t('hq_bulk_fail')} {failed}{t('rows')}", "warning")
        if not sent and not failed and not skipped:
            flash(t("bulk_remind_none"), "warning")
        return redirect(url_for("invoices", **request.args))

    @app.route("/invoices/bulk-issue", methods=["POST"])
    @admin_required
    def invoices_bulk_issue():
        db = Session()
        ids = parse_selected_invoice_ids(request.form)
        if not ids:
            flash(t("bulk_issue_none"), "warning")
            return redirect(url_for("invoices", **request.args))
        rows = (
            db.query(Invoice)
            .filter(Invoice.id.in_(ids), Invoice.status == "draft")
            .order_by(Invoice.year.asc(), Invoice.half.asc(), Invoice.seq.asc())
            .all()
        )
        if not rows:
            flash(t("bulk_issue_none"), "warning")
            return redirect(url_for("invoices", **request.args))
        if request.form.get("compose_ready") != "1":
            for row in rows:
                apply_training_headcounts(db, row)
            db.commit()
            return render_mail_compose(
                "issue",
                rows,
                url_for("invoices_bulk_issue", **request.args),
                url_for("invoices", **request.args),
                fill=False,
            )
        subject = request.form.get("mail_subject")
        body_text = request.form.get("mail_body")
        sent = failed = skipped = 0
        via = "outbox"
        for row in rows:
            to_list, cc_list = parse_compose_recipients(request.form, row, "issue")
            if not to_list:
                skipped += 1
                continue
            try:
                apply_training_headcounts(db, row)
                snapshot_invoice_party(row)
                path = rebuild_pdf(row)
                mail = send_invoice_email(
                    row,
                    path,
                    kind="issue",
                    subject=subject,
                    body_text=body_text,
                    to_recipients=to_list,
                    cc_recipients=cc_list,
                )
                if mail.get("ok"):
                    sent += 1
                    if mail.get("via") == "smtp":
                        via = "smtp"
                    row.status = "issued"
                    row.issued_at = utcnow()
                    row.issued_by = current_user().username
                    row.document_date = date.today()
                    log_action(db, "invoice", row.invoice_no, "ISSUE", mail_log_detail(mail))
                else:
                    failed += 1
            except Exception:
                failed += 1
        db.commit()
        if sent:
            flash(f"{t('bulk_issue_ok')} {sent}{t('rows')}", "success")
            if via == "outbox":
                flash(t("email_saved_outbox"), "info")
        if skipped:
            flash(f"{t('no_email')} {skipped}{t('rows')}", "warning")
        if failed:
            flash(f"{t('hq_bulk_fail')} {failed}{t('rows')}", "warning")
        if not sent and not failed and not skipped:
            flash(t("bulk_issue_none"), "warning")
        return redirect(url_for("invoices", **request.args))

    @app.route("/invoices/bulk-delete", methods=["POST"])
    @admin_required
    def invoices_bulk_delete():
        db = Session()
        ids = parse_selected_invoice_ids(request.form)
        if not ids:
            flash(t("bulk_delete_none"), "warning")
            return redirect(url_for("invoices", **request.args))
        rows = (
            db.query(Invoice)
            .filter(Invoice.id.in_(ids))
            .order_by(Invoice.year.asc(), Invoice.half.asc(), Invoice.seq.asc())
            .all()
        )
        if not rows:
            flash(t("bulk_delete_none"), "warning")
            return redirect(url_for("invoices", **request.args))
        deleted = 0
        paths = []
        for row in rows:
            paths.extend(delete_invoice_row(db, row))
            deleted += 1
        db.commit()
        unlink_paths(paths)
        flash(f"{t('bulk_delete_ok')} {deleted}{t('rows')}", "success")
        return redirect(url_for("invoices", **request.args))

    @app.route("/invoices/hq-copy", methods=["POST"])
    @admin_required
    def invoices_hq_copy():
        db = Session()
        if not hq_copy_emails():
            flash(t("hq_bulk_none"), "danger")
            return redirect(url_for("invoices", **request.args))
        rows = (
            filtered_invoices(db, current_user())
            .filter(Invoice.status != "draft")
            .order_by(Invoice.year.asc(), Invoice.half.asc(), Invoice.seq.asc())
            .all()
        )
        if not rows:
            flash(t("hq_bulk_none"), "warning")
            return redirect(url_for("invoices", **request.args))
        if request.form.get("compose_ready") != "1":
            return render_mail_compose(
                "hq",
                rows,
                url_for("invoices_hq_copy", **request.args),
                url_for("invoices", **request.args),
                fill=False,
            )
        subject = request.form.get("mail_subject")
        body_text = request.form.get("mail_body")
        sent = failed = 0
        via = "outbox"
        for row in rows:
            try:
                path = rebuild_pdf(row)
                mail = send_invoice_email(
                    row,
                    path,
                    kind="hq",
                    hq_only=True,
                    subject=subject,
                    body_text=body_text,
                )
                if mail.get("ok"):
                    sent += 1
                    if mail.get("via") == "smtp":
                        via = "smtp"
                    log_action(
                        db,
                        "invoice",
                        row.invoice_no,
                        "HQ_COPY",
                        ",".join(mail.get("recipients") or []),
                    )
                else:
                    failed += 1
            except Exception:
                failed += 1
        db.commit()
        if sent:
            flash(f"{t('hq_bulk_ok')} {sent}{t('rows')} · {hq_copy_emails()[0]}", "success")
            if via == "outbox":
                flash(t("email_saved_outbox"), "info")
        if failed:
            flash(f"{t('hq_bulk_fail')} {failed}{t('rows')}", "warning")
        if not sent and not failed:
            flash(t("hq_bulk_none"), "warning")
        return redirect(url_for("invoices", **request.args))

    @app.route("/invoices/<int:iid>/pdf")
    @login_required
    def invoice_pdf(iid):
        db = Session()
        user = current_user()
        row = db.get(Invoice, iid)
        if not row:
            return redirect(url_for("invoices"))
        if user.role != "admin" and row.subsidiary_id != user.subsidiary_id:
            flash(t("no_permission"), "danger")
            return redirect(url_for("invoices"))
        path = rebuild_pdf(row)
        db.commit()
        return send_file(path, as_attachment=True, download_name=f"{row.invoice_no}.pdf")

    @app.route("/invoices/template")
    @admin_required
    def invoice_template():
        wb = Workbook()
        ws = wb.active
        ws.title = "invoices"
        ws.append(["교육명", "교육일정", "교육담당자명", "법인코드", "학습자 이름", "원화 교육비"])
        ws.append(
            [
                "POSCO Group Leadership Program",
                "2026-03-02 ~ 2026-03-06",
                "Kim Minjun",
                "02VN01",
                "Nguyen Van A",
                1500000,
            ]
        )
        ws.append(
            [
                "POSCO Group Leadership Program",
                "2026-03-02 ~ 2026-03-06",
                "Kim Minjun",
                "02VN01",
                "Tran Thi B",
                1500000,
            ]
        )
        bio = BytesIO()
        wb.save(bio)
        bio.seek(0)
        return send_file(
            bio,
            as_attachment=True,
            download_name="invoice_upload_template.xlsx",
            mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        )

    @app.route("/invoices/upload", methods=["POST"])
    @admin_required
    def invoice_upload():
        db = Session()
        f = request.files.get("file")
        if not f or not allowed_file(f.filename):
            flash(t("invalid_file"), "danger")
            return redirect(url_for("invoices"))
        currency = (request.form.get("currency") or "USD").upper()
        if currency not in ("USD", "EUR"):
            currency = "USD"
        rate = _as_rate(request.form.get("exchange_rate") or 0)
        rate_date = _as_date(request.form.get("rate_date") or date.today())
        if rate <= 0:
            flash(t("required"), "danger")
            return redirect(url_for("invoices"))
        try:
            rows = load_table(f)
        except Exception:
            flash(t("upload_fail"), "danger")
            return redirect(url_for("invoices"))

        groups = {}
        errors = 0
        for row in rows:
            try:
                code = str(norm_key(row, "법인코드", "subsidiary_code", "code")).strip().upper()
                title = str(norm_key(row, "교육명", "education_name", "title")).strip()
                instructor = str(norm_key(row, "교육담당자명", "instructor")).strip()
                learner = str(norm_key(row, "학습자 이름", "learner_name", "learner")).strip()
                unit = _as_int(norm_key(row, "원화 교육비", "unit_price_krw", "amount_krw"))
                period_raw = norm_key(row, "교육일정", "period")
                start = _as_date(norm_key(row, "period_start", "start_date")) if norm_key(row, "period_start", "start_date") else None
                end = _as_date(norm_key(row, "period_end", "end_date")) if norm_key(row, "period_end", "end_date") else None
                ps, pe = parse_period(period_raw, start, end)
                if not code or not title or unit <= 0:
                    errors += 1
                    continue
                key = (code, title, ps.isoformat(), pe.isoformat())
                groups.setdefault(
                    key,
                    {"instructor": instructor, "items": []},
                )
                if instructor:
                    groups[key]["instructor"] = instructor
                groups[key]["items"].append({"learner": learner, "unit": unit})
            except Exception:
                errors += 1

        created = 0
        for (code, title, ps, pe), payload in groups.items():
            sub = db.query(Subsidiary).filter_by(code=code).first()
            if not sub:
                errors += 1
                continue
            period_start = _as_date(ps)
            period_end = _as_date(pe)
            year = period_start.year
            half = half_from_date(period_start)
            inv_no, seq = next_invoice_no(db, year, half)
            inv = Invoice(
                invoice_no=inv_no,
                year=year,
                half=half,
                seq=seq,
                subsidiary_id=sub.id,
                education_name=title,
                period_start=period_start,
                period_end=period_end,
                instructor=payload["instructor"],
                qty=0,
                unit_price_krw=0,
                amount_krw=0,
                currency=currency,
                exchange_rate=rate,
                rate_date=rate_date,
                amount_fx=0,
                vat_rate=0,
                status="draft",
                document_date=date.today(),
            )
            for item in payload["items"]:
                unit = item["unit"]
                inv.items.append(
                    InvoiceItem(
                        learner_name=item["learner"],
                        description=title,
                        qty=1,
                        unit_price_krw=unit,
                        amount_krw=unit,
                    )
                )
            apply_invoice_totals(inv)
            db.add(inv)
            db.flush()
            log_action(db, "invoice", inv.invoice_no, "BULK_CREATE", f"{code} {title}")
            created += 1
        db.commit()
        flash(f"{t('upload_ok')} · {t('created_n')} {created} / {t('errors')} {errors}", "success")
        return redirect(url_for("invoices"))

    @app.route("/collections")
    @admin_required
    def collections():
        return redirect(url_for("invoices"))

    @app.route("/logs")
    @admin_required
    def logs():
        db = Session()
        rows = db.query(AuditLog).order_by(AuditLog.created_at.desc()).limit(300).all()
        return render_template("logs.html", rows=rows)

    return app


app = create_app()

if __name__ == "__main__":
    app.run(host="0.0.0.0", port=5000, debug=True, use_reloader=False)
